#!/usr/bin/env python
"""Write Markdown, DOCX, and PDF summary reports."""

from __future__ import annotations

import importlib.util

import pandas as pd

from common import FIGURES, REPORTS, TABLES, ensure_dirs

SECTIONS = [
    (
        "Introduction",
        "This report documents a reproducible hierarchical mechanism-based machine-learning pipeline for explaining and predicting particulate organic carbon radiocarbon (Δ14C-POC) in six large Arctic rivers.",
    ),
    (
        "Data",
        "Inputs include ArcticGRO sample chemistry, river-specific daily discharge, HydroATLAS catchment attributes, REAL migration summaries, and reference literature. The cleaned modeling table is written to outputs/tables/cleaned_modeling_table.csv.",
    ),
    (
        "Methods",
        "The workflow constructs daily discharge climatologies, Q*, dQ/dt, rising/falling limb indicators, six hydrograph stages, thaw-entrainment bank-access indices, surface-erosion/runoff indices, and a catchment old-carbon baseline index.",
    ),
    (
        "Predictor screening",
        "The predictor audit removes δ13C and radiocarbon response leakage, raw/log duplicates, deterministic concentration or flux products, identifiers, and highly collinear predictors. It also flags double-counting risk.",
    ),
    (
        "Model construction",
        "Models include interpretable multiple-variable linear regression, polynomial Ridge regression, penalized spline additive Ridge, Random Forest, ExtraTrees, Gradient Boosting, and a hierarchical Ridge model. The hierarchical model decomposes Δ14C into catchment baseline, seasonal/daily tuning, and river residual calibration components.",
    ),
    (
        "Cross-validation design",
        "Leave-one-river-out validation estimates transferability to a new or unmonitored river. Blocked river-year validation estimates operational skill for monitored rivers with prior samples. Random row-level splits and in-sample performance are not used as main evidence.",
    ),
    (
        "Results",
        "Metrics are reported for R², RMSE, MAE, bias, median error, residual standard deviation, p95 absolute error, overprediction rate, underprediction rate, calibration slope/intercept, Pearson r, and Spearman ρ.",
    ),
    (
        "River-specific diagnostics",
        "Measured-vs-predicted and residual figures are generated separately by river for Ob, Yenisey, Lena, Kolyma, Yukon, and Mackenzie when validation predictions are available.",
    ),
    (
        "Mechanistic interpretation",
        "The interpretation distinguishes two controls: (1) a catchment template involving permafrost, lithology, wetland extent, migration potential, soil organic carbon, and erosion susceptibility; and (2) daily/seasonal source access involving discharge, temperature, TSS, POC%, hydrograph limb, bank-access gate, and surface-runoff gate.",
    ),
    (
        "Limitations",
        "REAL reach records are not full geometries, so migration matching is audited and may fall back to basin identifiers or a documented global summary. The six-river catchment table is small, so catchment-only models risk overfitting and should be interpreted as structured baselines.",
    ),
    (
        "Recommendations",
        "Use leave-one-river-out results for new-river transfer claims, blocked river-year results for monitored-river prediction claims, hierarchical decomposition for mechanistic interpretation, and the empirical equation only as a transparent approximation to nonlinear models.",
    ),
]


def metrics_summary(metrics: pd.DataFrame) -> pd.DataFrame:
    if metrics.empty or "rmse" not in metrics:
        return pd.DataFrame()
    return (
        metrics[metrics["rmse"].notna()]
        .groupby(["validation_scheme", "model_tier", "algorithm"], as_index=False)
        .agg(
            rmse=("rmse", "mean"),
            mae=("mae", "mean"),
            r2=("r2", "mean"),
            bias=("bias", "mean"),
            n=("n", "sum"),
        )
        .sort_values(["validation_scheme", "rmse"])
        .head(30)
    )


def markdown_report(metrics: pd.DataFrame) -> str:
    top = metrics_summary(metrics)
    text = ["# ArcticGRO Δ14C-POC prediction report", ""]
    for title, body in SECTIONS:
        text += [f"## {title}", body, ""]
        if title == "Results" and not top.empty:
            text += [
                "### Top validation metric summaries",
                top.to_markdown(index=False),
                "",
            ]
    text += ["## Tables", ""]
    for table in sorted(TABLES.glob("*.csv")):
        text.append(f"- `{table.relative_to(TABLES.parents[1])}`")
    text += ["", "## Figures", ""]
    for fig in sorted(FIGURES.glob("*.png")):
        text.append(f"- `{fig.relative_to(FIGURES.parents[1])}`")
    text += [""]
    return "\n".join(text)


def write_docx(text: str) -> None:
    if importlib.util.find_spec("docx") is None:
        (REPORTS / "final_ml_prediction_report.docx.txt").write_text(text)
        return
    from docx import Document

    doc = Document()
    for line in text.splitlines():
        if line.startswith("# "):
            doc.add_heading(line[2:], level=0)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=2)
        elif line.strip() and not line.startswith("|"):
            doc.add_paragraph(line)
    doc.add_heading("Selected figures", level=1)
    for fig in sorted(FIGURES.glob("*.png")):
        doc.add_heading(fig.stem.replace("_", " "), level=2)
        doc.add_picture(str(fig))
    doc.save(REPORTS / "final_ml_prediction_report.docx")


def write_pdf(text: str) -> None:
    if importlib.util.find_spec("reportlab") is None:
        (REPORTS / "final_ml_prediction_report.pdf.txt").write_text(text)
        return
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    pdf = canvas.Canvas(
        str(REPORTS / "final_ml_prediction_report.pdf"), pagesize=letter
    )
    _, height = letter
    y = height - 50
    for line in text.splitlines():
        if y < 50:
            pdf.showPage()
            y = height - 50
        pdf.setFont(
            "Helvetica-Bold" if line.startswith("#") else "Helvetica",
            12 if line.startswith("#") else 9,
        )
        pdf.drawString(40, y, line[:112])
        y -= 14
    pdf.save()


def main() -> None:
    ensure_dirs()
    metrics = (
        pd.read_csv(TABLES / "model_metrics.csv")
        if (TABLES / "model_metrics.csv").exists()
        else pd.DataFrame()
    )
    text = markdown_report(metrics)
    (REPORTS / "final_ml_prediction_report.md").write_text(text)
    write_docx(text)
    write_pdf(text)
    print(f"Wrote reports to {REPORTS}")


if __name__ == "__main__":
    main()
